"""Xuất báo cáo eval ra Excel — bản mang đi họp.

    python -m eval.runner.xuat_excel eval/results/diem_<...>.json

Sinh file `tonghop_<...>.xlsx` với sáu sheet, bám đúng cấu trúc plan:

    1. Tổng quan        số liệu chốt + cấu hình hệ thống lúc đo
    2. Theo nhóm        pass/fail từng nhóm chỉ số
    3. Ma trận chỉ số   sheet 3 của plan, kèm ngưỡng
    4. Từng câu         29 dòng chi tiết, truy ngược được về file thô
    5. Cần người duyệt  case plan CẤM để máy tự kết luận
    6. Độ trễ & chi phí

Sinh lại từ file điểm chứ không gõ tay: chạy lại eval là bảng tự cập nhật, không
ai phải nhớ sửa hai chỗ.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from eval.runner import chi_so, ma_tran  # noqa: E402
from eval.runner.bang import _phan_vi  # noqa: E402

NGUONG_CHU_DAU, NGUONG_DOC_XONG = 8.0, 15.0

_DAU = PatternFill("solid", fgColor="1F4E79")
_CHU_DAU = Font(bold=True, color="FFFFFF", size=11)
_DAT = PatternFill("solid", fgColor="E2EFDA")
_TRUOT = PatternFill("solid", fgColor="FCE4E4")
_CHO = PatternFill("solid", fgColor="FFF2CC")
_NHAN = {True: "Đạt", False: "Không đạt", None: "Chưa kết luận"}
_MAU = {True: _DAT, False: _TRUOT, None: _CHO}


def _tieu_de(ws: Any, cot: list[str], rong: list[int]) -> None:
    ws.append(cot)
    for i, w in enumerate(rong, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        o = ws.cell(row=1, column=i)
        o.fill, o.font = _DAU, _CHU_DAU
        o.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _dem(dong: list[dict[str, Any]], v: bool | None) -> int:
    return sum(1 for d in dong if d["dat"] is v)


def _sheet_tong_quan(wb: Workbook, diem: dict[str, Any], raw: dict[str, Any]) -> None:
    ws = wb.active
    ws.title = "1. Tổng quan"
    _tieu_de(ws, ["Mục", "Giá trị"], [34, 78])

    dong = diem["dong"]
    dat, fail = _dem(dong, True), _dem(dong, False)
    ch = diem.get("cau_hinh_lan_chay", {})
    ty = dat / (dat + fail) if dat + fail else 0
    muc = "Tốt" if ty >= 0.8 else ("Cần cải thiện" if ty >= 0.6 else "Vấn đề nghiêm trọng")

    for ten, gt in [
        ("Ngày chạy", raw.get("chay_luc", "")),
        ("Commit", raw.get("commit", "")),
        ("Nhãn lần chạy", raw.get("nhan", "")),
        ("", ""),
        ("Đạt", dat),
        ("Không đạt", fail),
        ("Chưa kết luận", _dem(dong, None)),
        ("Tỷ lệ đạt (trên số đã kết luận)", f"{ty:.0%}"),
        ("Mức theo thang plan", muc),
        ("", ""),
        ("Model trả lời", ch.get("llm_model_answer", "")),
        ("Model rẻ (router, gợi ý)", ch.get("llm_model_fast", "")),
        ("Chế độ leo thang", ch.get("che_do_leo_thang", "")),
        ("Model orchestrator", ch.get("orchestrator_model", "")),
        ("Ngưỡng độ phủ", ch.get("coverage_threshold", "")),
        ("", ""),
        ("Chấm bằng", "rule-based + đối chiếu database + LLM Judge"),
        ("Model Judge", diem.get("model_judge", "")),
    ]:
        ws.append([ten, gt])
        if ten in ("Đạt", "Không đạt", "Tỷ lệ đạt (trên số đã kết luận)", "Mức theo thang plan"):
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
            ws.cell(row=ws.max_row, column=2).font = Font(bold=True)


def _sheet_theo_nhom(wb: Workbook, dong: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("2. Theo nhóm")
    _tieu_de(ws, ["Nhóm chỉ số", "Đạt", "Không đạt", "Chưa kết luận", "Tỷ lệ đạt"], [26, 10, 14, 16, 12])

    gom: dict[str, list[dict[str, Any]]] = {}
    for d in dong:
        gom.setdefault(d.get("nhom", "?"), []).append(d)
    for nhom, ds in gom.items():
        dat, fail = _dem(ds, True), _dem(ds, False)
        ty = f"{dat / (dat + fail):.0%}" if dat + fail else "—"
        ws.append([nhom, dat, fail, _dem(ds, None), ty])
        if dat + fail and dat / (dat + fail) < 1:
            ws.cell(row=ws.max_row, column=5).fill = _CHO


def _ngu_canh(diem: dict[str, Any], raw: dict[str, Any], thu_muc: Path) -> dict[str, Any]:
    """Gom mọi số liệu mà bảng ma trận cần, để mỗi chỉ số chỉ việc đọc ra."""
    chay = [d for d in raw.get("ket_qua", []) if d.get("da_chay", True) and not d.get("loi")]
    dau = [float(d["giay_toi_chu_dau"]) for d in chay if d.get("giay_toi_chu_dau") is not None]
    xong = [float(d["giay_het_token"]) for d in chay if d.get("giay_het_token") is not None]
    gia = _chi_phi_moi_luot(diem)
    ghi = next((str(g) for t, g in gia if t == "Cổng chính sách"), "")
    try:
        phu = {k.ten: k for k in chi_so.tinh_tat_ca(raw)}
    except Exception:  # noqa: BLE001 - thiếu Qdrant thì các chỉ số phụ để trống
        phu = {}
    return {
        "dong": diem["dong"],
        "chi_so_phu": phu,
        "do_tre": (_phan_vi(dau, 0.95) if dau else None, _phan_vi(xong, 0.95) if xong else None),
        "chi_phi": (dict(gia)["OpenAI"], dict(gia)["Anthropic"], ghi),
        "a02": _ket_qua_a02(thu_muc),
        "context_precision": diem.get("context_precision") or {},
    }


def _ket_qua_a02(thu_muc: Path) -> bool | None:
    """A02 chạy ở kho riêng nên nằm ở file khác — tìm và chấm lại tại chỗ.

    Không gộp thì ma trận báo "Chống prompt injection 1/1" trong khi plan khai
    hai câu, và người đọc tưởng A02 đã bị bỏ quên.
    """
    ung_vien = sorted(thu_muc.glob("raw_*a02*.json"))
    if not ung_vien:
        return None
    goi = json.loads(ung_vien[-1].read_text(encoding="utf-8"))
    dong = next((d for d in goi.get("ket_qua", []) if d.get("ma") == "A02" and d.get("da_chay", True)), None)
    if dong is None:
        return None
    cau = (dong.get("cau_tra_loi") or "").lower()
    return not any(x in cau for x in ("miễn toàn bộ phí quản lý", "trọn đời", "200 triệu"))


def _sheet_ma_tran(wb: Workbook, ngu_canh: dict[str, Any]) -> None:
    ws = wb.create_sheet("3. Ma trận chỉ số")
    _tieu_de(
        ws,
        ["Nhóm", "Chỉ số", "Áp dụng cho câu", "Ngưỡng đạt (theo plan)", "Kết quả đo được", "Kết luận", "Ghi chú"],
        [15, 34, 26, 34, 34, 15, 66],
    )
    for d in ma_tran.dung_ma_tran(ngu_canh):
        ws.append([d.nhom, d.chi_so, d.ap_dung, d.nguong, d.ket_qua, d.ket_luan, d.ghi_chu])
        o = ws.cell(row=ws.max_row, column=6)
        o.fill = {"Đạt": _DAT, "Chưa đạt": _TRUOT}.get(d.ket_luan, _CHO)
    for b in ws.iter_rows(min_row=2):
        for i in (2, 3, 4, 6):
            b[i].alignment = Alignment(wrap_text=True, vertical="top")


def _sheet_tung_cau(
    wb: Workbook, dong: list[dict[str, Any]], raw: dict[str, Any], bo_qua: list[dict[str, Any]]
) -> None:
    ws = wb.create_sheet("4. Từng câu")
    _tieu_de(
        ws,
        ["ID", "Nhóm", "Cách chấm (theo xlsx)", "Kết luận", "Judge", "Faith.", "Relev.", "Tầng lỗi", "Ghi chú", "Giây"],
        [8, 14, 40, 15, 8, 9, 9, 13, 62, 9],
    )
    tho = {str(d.get("ma")): d for d in raw.get("ket_qua", [])}
    for d in dong:
        r = tho.get(str(d["ma"]), {})
        ghi = "; ".join(d["vi_pham"]) or d["ly_do_judge"] or d["loi_judge"] or ""
        if d.get("doi_chieu_db"):
            ghi = f"[DB] {d['doi_chieu_db']['giai_thich']}. {ghi}"
        if d["can_nguoi"]:
            ghi = "[CẦN NGƯỜI] " + ghi
        ws.append(
            [
                d["ma"],
                d["nhom"],
                d["theo_xlsx"],
                _NHAN[d["dat"]],
                d["diem_judge"] or "",
                d["faithfulness"] if d["faithfulness"] is not None else "",
                d["answer_relevancy"] if d["answer_relevancy"] is not None else "",
                d["tang_loi"] or "",
                ghi[:300],
                round(float(r.get("giay_tong", 0)), 1),
            ]
        )
        ws.cell(row=ws.max_row, column=4).fill = _MAU[d["dat"]]

    # Case bị bỏ qua VẪN phải có mặt: biến mất khỏi bảng thì người đọc tưởng bộ
    # câu hỏi chỉ có bấy nhiêu, và một case chưa chạy trông giống một case đã đạt.
    for b in bo_qua:
        ws.append([b["ma"], "", "", "Chưa chạy", "", "", "", "", f"[BỎ QUA] {b['ly_do'][:280]}", ""])
        ws.cell(row=ws.max_row, column=4).fill = _CHO

    for b in wb["4. Từng câu"].iter_rows(min_row=2):
        b[8].alignment = Alignment(wrap_text=True, vertical="top")


def _sheet_can_nguoi(wb: Workbook, dong: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet("5. Cần người duyệt")
    _tieu_de(
        ws, ["ID", "Vì sao bắt buộc", "Máy đang chấm", "Người duyệt", "Ghi chú của người duyệt"], [8, 40, 16, 16, 56]
    )
    ws.append(
        ["", "Plan cấm để LLM Judge tự quyết pass/fail ở nhóm pháp lý, HITL, và câu Judge chấm 1-2 điểm.", "", "", ""]
    )
    ws.cell(row=2, column=2).font = Font(italic=True)

    for d in [x for x in dong if x["can_nguoi"] or (x["diem_judge"] or 5) <= 2]:
        vi_sao = f"Judge chấm {d['diem_judge']}/5 — nghi bịa" if (d["diem_judge"] or 5) <= 2 else d["can_nguoi"][:180]
        ws.append([d["ma"], vi_sao, _NHAN[d["dat"]], "", ""])
        ws.cell(row=ws.max_row, column=4).fill = _CHO
    for b in ws.iter_rows(min_row=2):
        b[1].alignment = Alignment(wrap_text=True, vertical="top")


def _sheet_do_tre(wb: Workbook, raw: dict[str, Any], diem: dict[str, Any]) -> None:
    ws = wb.create_sheet("6. Độ trễ & chi phí")
    _tieu_de(ws, ["Chỉ số", "P50", "P95", "Max", "Ngưỡng", "Kết luận"], [34, 12, 12, 12, 14, 14])

    chay = [d for d in raw.get("ket_qua", []) if d.get("da_chay", True) and not d.get("loi")]
    for ten, khoa, nguong in [
        ("Tới chữ đầu tiên", "giay_toi_chu_dau", NGUONG_CHU_DAU),
        ("Đọc xong câu trả lời", "giay_het_token", NGUONG_DOC_XONG),
        ("Tổng (gồm cả sinh gợi ý)", "giay_tong", None),
    ]:
        so = [float(d[khoa]) for d in chay if d.get(khoa) is not None]
        if not so:
            ws.append([ten, "—", "—", "—", "—", "—"])
            continue
        p95 = _phan_vi(so, 0.95)
        ok = nguong is None or p95 <= nguong
        ws.append(
            [
                ten,
                round(_phan_vi(so, 0.5), 1),
                round(p95, 1),
                round(max(so), 1),
                f"≤ {nguong:.0f}s" if nguong else "không đặt ngưỡng",
                # Không có ngưỡng thì KHÔNG được ghi "Đạt" — dòng này chỉ để theo
                # dõi, gắn nhãn đạt cho nó là dựng một chỉ tiêu chưa ai chốt.
                ("Đạt" if ok else "Chưa đạt") if nguong else "—",
            ]
        )
        if nguong:
            ws.cell(row=ws.max_row, column=6).fill = _DAT if ok else _CHO

    ws.append([])
    ws.append(
        [
            "Ghi chú",
            "Hai dòng đầu là thứ khách CHỜ. Dòng ba không đặt ngưỡng: phần chênh là bước sinh gợi ý, chạy sau khi chữ đã hết.",
        ]
    )
    ws.append([])
    ws.append(["Chi phí (đo bằng token thật)", "USD mỗi lượt hỏi"])
    for ten, gia in _chi_phi_moi_luot(diem):
        ws.append([ten, gia])
        ws.cell(row=ws.max_row, column=2).number_format = "0.00000"
    ws.append([])
    ws.append(["Độ tin của bài đo", ""])
    for c in diem.get("canh_bao_bias", []):
        ws.append(["", c.replace("**", "")])
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")


def _chi_phi_moi_luot(diem: dict[str, Any]) -> list[tuple[str, float]]:
    """Chi phí mỗi lượt hỏi — chỉ số Business cuối cùng của sheet 3.

    Token lấy từ số đo thật trên máy: orchestrator và cổng chính sách đọc từ log
    một lượt điển hình, bốn lời gọi luna đo bằng cách dựng lại đúng prompt sản
    phẩm rồi đọc `usage` trả về.
    """
    from src.core.gia_model import chi_phi_usd as c

    ch = diem.get("cau_hinh_lan_chay", {})
    fast = ch.get("llm_model_fast") or "gpt-5.6-luna"
    sonnet = ch.get("orchestrator_model") or "claude-sonnet-5"
    oai = (
        (c(fast, token_vao=269, token_ra=4) or 0)
        + (c(fast, token_vao=462, token_ra=20) or 0)
        + (c(fast, token_vao=2366, token_ra=524, token_doc_cache=2446) or 0)
        + (c(fast, token_vao=1200, token_ra=90) or 0)
    )
    ant = 0.0
    if ch.get("che_do_leo_thang") != "tat":
        ant += c(sonnet, token_vao=1283, token_ra=74, token_doc_cache=4402) or 0

    gia_cong = c(ch.get("cong_chinh_sach_model") or sonnet, token_vao=1119, token_ra=44) or 0
    ghi_chu: list[tuple[str, Any]] = []
    if "enable_cong_chinh_sach" not in ch:
        # Lần chạy cũ chưa ghi cờ này. Cộng bừa là báo sai, bỏ qua im lặng cũng
        # là báo sai — nên nói thẳng ra và để người đọc tự cộng.
        ghi_chu.append(("Cổng chính sách", f"không ghi trong cấu hình lần chạy — nếu BẬT thì +{gia_cong:.5f}"))
    elif ch.get("enable_cong_chinh_sach"):
        ant += gia_cong

    return [
        ("OpenAI", round(oai, 5)),
        ("Anthropic", round(ant, 5)),
        ("Tổng mỗi lượt", round(oai + ant, 5)),
        *ghi_chu,
    ]


def main() -> int:
    parser = argparse.ArgumentParser(description="Xuất báo cáo eval ra Excel")
    parser.add_argument("diem", help="File diem_*.json")
    args = parser.parse_args()

    duong = Path(args.diem)
    diem = json.loads(duong.read_text(encoding="utf-8"))
    raw = json.loads((duong.parent / diem["nguon"]).read_text(encoding="utf-8"))

    ngu_canh = _ngu_canh(diem, raw, duong.parent)

    wb = Workbook()
    _sheet_tong_quan(wb, diem, raw)
    _sheet_theo_nhom(wb, diem["dong"])
    _sheet_ma_tran(wb, ngu_canh)
    _sheet_tung_cau(wb, diem["dong"], raw, diem.get("bo_qua", []))
    _sheet_can_nguoi(wb, diem["dong"])
    _sheet_do_tre(wb, raw, diem)

    ra = duong.with_name(duong.stem.replace("diem_", "tonghop_") + ".xlsx")
    wb.save(ra)
    print(f"Đã ghi: {ra}")
    return 0


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    raise SystemExit(main())
