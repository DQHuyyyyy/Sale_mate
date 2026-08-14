"""Đọc kế hoạch crawl từ `Link_data.xlsx`.

File này là NGUỒN SỰ THẬT cho việc cào dữ liệu: thêm một dòng vào Excel là thêm
một tài liệu, không phải sửa code.

Bốn cột:
    Khu Vực          dự án, ví dụ "Vinhom OceanPark 2" — có thể để trống
    Nhóm thông tin   tên tài liệu, cũng là TÊN FILE
    Lấy thông tin gì hướng dẫn trích xuất, hoặc chính nội dung nếu đã có sẵn
    Link             URL nguồn

Ô gộp phải đọc theo VÙNG GỘP THẬT, không suy diễn "trống thì lấy dòng trên".
Vùng `A9:A12` chỉ phủ tới Ocean Park 3, nên hai dòng cuối (chính sách lãi suất,
hệ thống tiện ích) không thuộc dự án nào — suy diễn kế thừa sẽ gắn nhầm chúng
vào Ocean Park 3 rồi lọc sai về sau.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_PLAN_FILE = REPO_ROOT / "src" / "data" / "Link_data.xlsx"

_COT = {"khu_vuc": 0, "nhom": 1, "lay_gi": 2, "link": 3}

# Nội dung đã điền sẵn trong cột "Lấy thông tin gì" thì dài hơn hẳn một câu
# hướng dẫn. Trên ngưỡng này coi là nội dung, không phải chỉ dẫn.
_NGUONG_NOI_DUNG = 300


@dataclass
class TaiLieuCanCrawl:
    """Một tài liệu cần tạo ra, gom từ một hoặc nhiều dòng Excel."""

    ten: str  # = cột "Nhóm thông tin", dùng làm tên file
    khu_vuc: str
    link: str
    muc_can_lay: list[str] = field(default_factory=list)
    noi_dung_co_san: str = ""

    @property
    def can_crawl(self) -> bool:
        """Đã có nội dung sẵn thì không cào nữa — người dùng tự điền vào Excel."""
        return not self.noi_dung_co_san

    def ten_file(self, trung_ten: set[str]) -> str:
        """Tên file lấy y cột 'Nhóm thông tin'.

        Chỉ thêm tiền tố khu vực cho tên bị trùng giữa hai dự án — "Ưu đãi" có
        ở cả Ocean Park 2 lẫn 3, để nguyên thì file này ghi đè file kia.
        """
        if self.ten in trung_ten and self.khu_vuc:
            return f"{self.khu_vuc} - {self.ten}"
        return self.ten


def _doc_o_gop(ws) -> dict[tuple[int, int], object]:
    """Trả về giá trị thật của mọi ô, kể cả ô nằm trong vùng gộp."""
    gia_tri: dict[tuple[int, int], object] = {}
    for row in ws.iter_rows():
        for cell in row:
            gia_tri[(cell.row, cell.column)] = cell.value

    for vung in ws.merged_cells.ranges:
        goc = gia_tri.get((vung.min_row, vung.min_col))
        for r in range(vung.min_row, vung.max_row + 1):
            for c in range(vung.min_col, vung.max_col + 1):
                gia_tri[(r, c)] = goc
    return gia_tri


def doc_ke_hoach(path: Path = DEFAULT_PLAN_FILE) -> list[TaiLieuCanCrawl]:
    """Đọc Excel, gom các dòng cùng (Nhóm thông tin, Khu Vực) thành một tài liệu."""
    if not path.is_file():
        raise FileNotFoundError(f"Không tìm thấy kế hoạch crawl: {path}")

    ws = openpyxl.load_workbook(path, data_only=True).active
    o = _doc_o_gop(ws)

    gom: dict[tuple[str, str], TaiLieuCanCrawl] = {}
    for r in range(2, ws.max_row + 1):
        lay = lambda ten: str(o.get((r, _COT[ten] + 1)) or "").strip()  # noqa: E731
        nhom, link = lay("nhom"), lay("link")
        if not nhom or not link:
            continue

        khoa = (nhom, lay("khu_vuc"))
        tai_lieu = gom.setdefault(khoa, TaiLieuCanCrawl(ten=nhom, khu_vuc=lay("khu_vuc"), link=link))

        lay_gi = lay("lay_gi")
        if not lay_gi:
            continue
        if len(lay_gi) >= _NGUONG_NOI_DUNG:
            tai_lieu.noi_dung_co_san = lay_gi
        else:
            tai_lieu.muc_can_lay.append(lay_gi)

    return list(gom.values())


def ten_bi_trung(tai_lieu: list[TaiLieuCanCrawl]) -> set[str]:
    """Tên xuất hiện ở nhiều khu vực khác nhau — cần thêm tiền tố để khỏi đè nhau."""
    theo_ten: dict[str, set[str]] = {}
    for t in tai_lieu:
        theo_ten.setdefault(t.ten, set()).add(t.khu_vuc)
    return {ten for ten, khu in theo_ten.items() if len(khu) > 1}
